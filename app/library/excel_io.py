"""素材登记表 Excel 导入导出。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..db import now
from ..state import Ctx

COLUMNS = [
    ("哈希", "sha256", False), ("相对路径", "rel_path", False), ("状态", "status", False),
    ("城市", "city", True), ("景点", "spot", True), ("类型", "type", True),
    ("服务内容", "service", True), ("能否公开使用", "publishable", True),
    ("评分", "rating", True), ("备注", "note", True), ("时长(秒)", "duration", False),
]
EDITABLE = {key for _h, key, editable in COLUMNS if editable}


def export_xlsx(ctx: Ctx, path: str | Path, material_types: list[str]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "素材登记表"
    ws.append([h for h, _k, _e in COLUMNS])
    gray = PatternFill("solid", fgColor="DDDDDD")
    for i, (_h, _k, editable) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        if not editable:
            cell.fill = gray
    rows = ctx.db.all("SELECT * FROM material ORDER BY rel_path")
    for r in rows:
        ws.append([round(r[k], 1) if k == "duration" else r[k] for _h, k, _e in COLUMNS])
    n = len(rows) + 1
    col = {k: i for i, (_h, k, _e) in enumerate(COLUMNS, start=1)}

    def letter(k):
        return ws.cell(row=1, column=col[k]).column_letter

    dv_pub = DataValidation(type="list", formula1='"是,否,未确认"', allow_blank=True)
    dv_type = DataValidation(type="list", formula1='"' + ",".join(material_types) + '"', allow_blank=True)
    dv_rate = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    for dv, k in ((dv_pub, "publishable"), (dv_type, "type"), (dv_rate, "rating")):
        ws.add_data_validation(dv)
        dv.add(f"{letter(k)}2:{letter(k)}{max(n, 2) + 500}")
    widths = {"sha256": 14, "rel_path": 46, "note": 30, "service": 18}
    for k, i in col.items():
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(k, 12)
    ws.freeze_panes = "A2"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_xlsx(ctx: Ctx, path: str | Path, material_types: list[str]) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    key_by_header = {h: k for h, k, _e in COLUMNS}
    idx = {key_by_header[h]: i for i, h in enumerate(header) if h in key_by_header}
    if "sha256" not in idx and "rel_path" not in idx:
        return "导入失败：表格中没有“哈希”或“相对路径”列。"
    updated, unmatched, problems = 0, 0, []
    for line_no, row in enumerate(rows, start=2):
        if row is None or all(v is None for v in row):
            continue

        def val(k):
            i = idx.get(k)
            v = row[i] if i is not None and i < len(row) else None
            return "" if v is None else str(v).strip()

        m = None
        if val("sha256"):
            m = ctx.db.one("SELECT id FROM material WHERE sha256 = ?", (val("sha256"),))
        if m is None and val("rel_path"):
            m = ctx.db.one("SELECT id FROM material WHERE rel_path = ?", (val("rel_path").replace("\\", "/"),))
        if m is None:
            unmatched += 1
            continue
        changes = {}
        for k in EDITABLE:
            if k not in idx:
                continue
            v = val(k)
            if k == "publishable":
                if v not in ("是", "否", "未确认"):
                    if v:
                        problems.append(f"第{line_no}行 能否公开使用 “{v}” 无效，已忽略")
                    continue
            if k == "type" and v and v not in material_types:
                problems.append(f"第{line_no}行 类型 “{v}” 不在类型列表中，已忽略")
                continue
            if k == "rating":
                try:
                    v = max(0, min(5, int(float(v or 0))))
                except ValueError:
                    problems.append(f"第{line_no}行 评分 “{v}” 无效，已忽略")
                    continue
            changes[k] = v
        if changes:
            changes["updated_at"] = now()
            ctx.db.update("material", m["id"], changes)
            updated += 1
    msg = f"导入完成：更新 {updated} 条"
    if unmatched:
        msg += f"，{unmatched} 行没有找到对应素材"
    msg += "。"
    if problems:
        msg += "\n" + "\n".join(problems[:30])
    return msg
